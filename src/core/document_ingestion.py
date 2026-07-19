"""
Canonical document ingestion and normalization.

Converts supported file types and raw text into a structured
single-page markdown representation so the rest of the system can
continue through the same assembler and validator path.
"""

from __future__ import annotations


import csv
import json
import sqlite3
from dataclasses import dataclass
from pathlib import Path

from core.logger import logger

from models.document import (
    DocumentMetadata,
    DocumentType,
    ExtractionResponse,
    PageExtraction,
    PageMetadata,
)


@dataclass(slots=True)
class IngestedDocument:
    """
    Normalized content ready for assembly.
    """

    metadata: DocumentMetadata

    pages: list[PageExtraction]

    source_path: Path | None = None

    source_kind: str = "unknown"


class DocumentIngestionService:
    """
    Normalizes non-PDF inputs into canonical page extractions.
    """

    TEXT_EXTENSIONS = {
        ".txt",
        ".md",
        ".sql",
        ".json",
        ".csv",
        ".tsv",
        ".log",
    }

    DATABASE_EXTENSIONS = {
        ".db",
        ".sqlite",
        ".sqlite3",
    }

    def ingest_file(self, file_path: Path) -> IngestedDocument:
        suffix = file_path.suffix.lower()

        if suffix == ".pdf":
            raise ValueError("PDF files are handled by the OCR pipeline")

        if suffix in self.TEXT_EXTENSIONS:
            content = self._read_text_file(file_path)
            source_kind = "text"
        elif suffix == ".docx":
            content = self._read_docx(file_path)
            source_kind = "docx"
        elif suffix == ".xlsx":
            content = self._read_xlsx(file_path)
            source_kind = "xlsx"
        elif suffix in self.DATABASE_EXTENSIONS:
            content = self._read_sqlite(file_path)
            source_kind = "sqlite"
        else:
            raise ValueError(f"Unsupported file type: {suffix}")

        return self._build_document(
            filename=file_path.name,
            source_kind=source_kind,
            content=content,
            file_size_bytes=file_path.stat().st_size,
        )

    def ingest_text(self, text: str, label: str = "legacy_text") -> IngestedDocument:
        return self._build_document(
            filename=f"{label}.txt",
            source_kind="legacy_text",
            content=text.strip(),
            file_size_bytes=len(text.encode("utf-8")),
        )

    def _read_text_file(self, file_path: Path) -> str:
        if file_path.suffix.lower() == ".json":
            try:
                data = json.loads(file_path.read_text(encoding="utf-8"))
                return json.dumps(data, indent=2, ensure_ascii=False)
            except json.JSONDecodeError:
                return file_path.read_text(encoding="utf-8")

        return file_path.read_text(encoding="utf-8", errors="replace")

    def _read_docx(self, file_path: Path) -> str:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise RuntimeError("python-docx is required for DOCX ingestion") from exc

        document = DocxDocument(file_path)

        blocks: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                blocks.append(text)

        for table_index, table in enumerate(document.tables, start=1):
            blocks.append(f"## Table {table_index}")
            for row in table.rows:
                blocks.append(" | ".join(cell.text.strip() for cell in row.cells))

        return "\n\n".join(blocks).strip()

    def _read_xlsx(self, file_path: Path) -> str:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("openpyxl is required for XLSX ingestion") from exc

        workbook = load_workbook(file_path, data_only=True, read_only=True)

        blocks: list[str] = []

        for sheet in workbook.worksheets:
            blocks.append(f"# Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if cell is None else str(cell) for cell in row]
                if any(value for value in values):
                    blocks.append(" | ".join(values))

        return "\n\n".join(blocks).strip()

    def _read_sqlite(self, file_path: Path) -> str:
        connection = sqlite3.connect(file_path)
        connection.row_factory = sqlite3.Row

        blocks: list[str] = []

        try:
            cursor = connection.execute(
                "SELECT name, sql FROM sqlite_master WHERE type='table' ORDER BY name"
            )

            for table in cursor.fetchall():
                table_name = table["name"]
                blocks.append(f"# Table: {table_name}")
                if table["sql"]:
                    blocks.append(f"```sql\n{table['sql']}\n```")

                sample_cursor = connection.execute(f'SELECT * FROM "{table_name}" LIMIT 20')
                rows = sample_cursor.fetchall()

                if rows:
                    headers = rows[0].keys()
                    blocks.append(" | ".join(headers))
                    blocks.append(" | ".join(["---"] * len(headers)))
                    for row in rows:
                        blocks.append(" | ".join("" if row[h] is None else str(row[h]) for h in headers))

        finally:
            connection.close()

        return "\n\n".join(blocks).strip()

    def _build_document(
        self,
        filename: str,
        source_kind: str,
        content: str,
        file_size_bytes: int,
    ) -> IngestedDocument:
        markdown = content.strip() or "[NO CONTENT]"

        metadata = DocumentMetadata(
            filename=filename,
            total_pages=1,
            document_type=DocumentType.UNKNOWN,
            encrypted=False,
            file_size_bytes=file_size_bytes,
            recommended_dpi=0,
            estimated_runtime_seconds=max(1.0, len(markdown) / 1000.0),
            pages=[
                PageMetadata(
                    page_number=1,
                    width=0,
                    height=0,
                    rotation=0,
                    dpi=0,
                    scanned=False,
                )
            ],
        )

        pages = [
            PageExtraction(
                page_number=1,
                regions=[
                    ExtractionResponse(
                        region_id=1,
                        markdown=markdown,
                        confidence=1.0,
                        tokens_used=max(1, len(markdown) // 4),
                        processing_time=0.0,
                    )
                ],
            )
        ]

        logger.info(
            "Ingested %s as %s document",
            filename,
            source_kind,
        )

        return IngestedDocument(
            metadata=metadata,
            pages=pages,
            source_path=None,
            source_kind=source_kind,
        )